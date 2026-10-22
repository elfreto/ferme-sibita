"""Convertit les fiches Excel de bandes en src/data/historique.json.

Usage : python3 -I scripts/import_excel.py fichier1.xlsx [fichier2.xlsx ...]
"""
import datetime, json, re, sys
import openpyxl

MOIS = ['janvier', 'février', 'mars', 'avril', 'mai', 'juin', 'juillet', 'août', 'septembre', 'octobre', 'novembre', 'décembre']

# Catégories du bilan Excel
CAT_KEYWORDS = [
    (4, ['salaire']),
    (2, ['véto', 'veto', 'phyto', 'médicament', 'chauffage', 'antitoxine', 'désinfect', 'calcium']),
    (3, ['provende', 'maïs', 'mais', 'soja', 'arachide', 'concentré', 'son de blé', 'coquille', 'huile', 'farine', 'ingrédient']),
    (1, ['poussin']),
]


def day(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    return None


def categorie(designation, pendant_vente):
    d = designation.lower()
    for cat, mots in CAT_KEYWORDS:
        if any(m in d for m in mots):
            return cat
    return 5 if pendant_vente else 2


def parse_notes_depenses(texte):
    """'500 fiches 500 taxi, 10000 planche' -> [(500,'fiches'), (500,'taxi'), (10000,'planche')]"""
    t = texte.split(':', 1)[1] if ':' in texte else texte
    t = re.sub(r'(\d) (\d{3})', r'\1\2', t)  # 10 000 -> 10000
    out = []
    for m in re.finditer(r'(\d+)\s*([^\d,]*)|([^\d,]+?)\s+(\d+)', t):
        if m.group(1):
            montant, lib = int(m.group(1)), m.group(2).strip(' ,')
        else:
            montant, lib = int(m.group(4)), m.group(3).strip(' ,')
        if montant > 0:
            out.append((montant, lib or 'divers'))
    return out


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    cell = lambda r, c: ws.cell(r, c).value
    bande = {'bilanExcel': {}}
    journal, ventes, depenses = [], [], []
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if v == 'Nombe de Poussins :' and 'poussins' not in bande:
                bande['poussins'] = cell(c.row, c.column + 2)
                bande['responsable'] = cell(c.row + 1, c.column + 2)
                bande['dateReception'] = day(cell(c.row + 2, c.column + 2))
            if isinstance(v, str) and v.startswith('Bilan'):
                for r in range(c.row + 1, c.row + 25):
                    lab = cell(r, c.column)
                    val = next((cell(r, c.column + k) for k in range(1, 4) if isinstance(cell(r, c.column + k), (int, float))), None)
                    if lab and val is not None:
                        bande['bilanExcel'][lab.strip()] = round(val, 4)
            if v != 'Date':
                continue
            if cell(c.row, c.column + 2) == 'Mortalité':
                for r in range(c.row + 1, ws.max_row + 1):
                    age = cell(r, c.column + 1)
                    if not isinstance(age, (int, float)):
                        break
                    morts = cell(r, c.column + 2)
                    if morts is None:
                        continue
                    journal.append({
                        'date': day(cell(r, c.column)), 'age': int(age), 'morts': int(morts),
                        'alimentKg': round(cell(r, c.column + 5) or 0, 2) or None,
                        'poidsMoyenKg': cell(r, c.column + 6), 'note': cell(r, c.column + 7),
                    })
            elif cell(c.row, c.column + 2) == 'Qté':
                premiere_date = derniere_date = None
                for r in range(c.row + 1, c.row + 6):
                    # Une ligne sans date est vendue le même jour que la ligne au-dessus
                    d = day(cell(r, c.column)) or derniere_date
                    qte, pu = cell(r, c.column + 2), cell(r, c.column + 3)
                    derniere_date = d
                    premiere_date = premiere_date or d
                    if qte:
                        ventes.append({'date': d, 'qte': int(qte), 'prixUnitaire': pu})
                    note = cell(r, c.column + 5)
                    if isinstance(note, str) and note.strip().lower().startswith('dépenses') and premiere_date:
                        for montant, lib in parse_notes_depenses(note):
                            depenses.append({'date': premiere_date, 'categorie': 5, 'designation': lib,
                                             'montant': montant, 'payeAvecVentes': True, 'origine': 'fiche de vente'})
    # Onglet Depenses : on s'arrête à la section qui finance la bande suivante
    if 'Depenses' in wb.sheetnames:
        pendant_vente = False
        lignes = list(wb['Depenses'].iter_rows(values_only=True))
        if any('argent vente' in ' '.join(str(v) for v in r if v).lower() and 'désignation' in ' '.join(str(v) for v in r if v).lower() for r in lignes):
            # Les petites dépenses de vente sont déjà regroupées dans l'onglet : on ignore le détail des fiches
            depenses[:] = [x for x in depenses if x['origine'] != 'fiche de vente']
        for row in lignes:
            vals = [v for v in row if v is not None]
            texte = ' '.join(str(v) for v in vals).lower()
            if 'prochaine bande' in texte:
                break
            if 'argent vente' in texte and 'désignation' in texte:
                pendant_vente = True
                continue
            dates = [day(v) for v in vals if isinstance(v, datetime.datetime)]
            nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            libs = [v for v in vals if isinstance(v, str) and v.lower() not in ('ok', 'argent vente')]
            if not nums or not libs or libs[0].lower().startswith('total'):
                continue
            avec_ventes = pendant_vente or 'argent vente' in texte
            depenses.append({'date': dates[0] if dates else bande['dateReception'],
                             'categorie': categorie(libs[0], avec_ventes), 'designation': libs[0],
                             'montant': nums[0], 'payeAvecVentes': avec_ventes, 'origine': 'onglet Dépenses'})
    d = datetime.date.fromisoformat(bande['dateReception'])
    bande['nom'] = f"Bande {MOIS[d.month - 1]} {d.year}"
    bande['id'] = f"import-{d.isoformat()}"
    for lst in (journal, ventes, depenses):
        for i, x in enumerate(lst):
            x['bandeId'] = bande['id']
            x['id'] = f"{bande['id']}-{'jvd'[[journal, ventes, depenses].index(lst)]}{i}"
    return bande, journal, ventes, depenses


if __name__ == '__main__':
    data = {'bandes': [], 'journal': [], 'ventes': [], 'depenses': []}
    for p in sys.argv[1:]:
        b, j, v, dep = extract(p)
        data['bandes'].append(b); data['journal'] += j; data['ventes'] += v; data['depenses'] += dep
    data['bandes'].sort(key=lambda b: b['dateReception'])
    # La dernière bande est en cours dans l'espace 1, les autres sont clôturées
    for b in data['bandes']:
        b['statut'] = 'cloturee'
        b['espace'] = 1
    data['bandes'][-1]['statut'] = 'en_cours'
    json.dump(data, open('src/data/historique.json', 'w'), ensure_ascii=False, indent=1)
    for b in data['bandes']:
        dep = [x for x in data['depenses'] if x['bandeId'] == b['id']]
        tot = {k: sum(x['montant'] for x in dep if x['categorie'] == k) for k in range(1, 6)}
        ven = [x for x in data['ventes'] if x['bandeId'] == b['id']]
        print(b['nom'], b['poussins'], 'jours', sum(1 for x in data['journal'] if x['bandeId'] == b['id']),
              'vendus', sum(x['qte'] for x in ven), 'CA', sum(x['qte'] * x['prixUnitaire'] for x in ven),
              'charges', tot, 'total', sum(tot.values()))
